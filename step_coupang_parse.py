"""
쿠팡이츠 정산내역서 파싱

쿠팡이츠는 정산서가 메일로 오지 않아 사이트에서 직접 다운로드한다.
  쿠팡이츠 로그인 > 매출관리 > 매출내역서 다운로드 > 매장선택(79대포 방이점) > 기간설정

손익 항목 ← 정산내역서 열 전체 합산 (사장님 지정):
  매출            = K열 (주문금액)
  광고비(중개)     = Q열
  배달비          = V열
  카드/어플수수료  = S열
  부가세          = AG열
  광고비(맞춤)     = AK열
  (식자재/부자재는 토더 쿠팡채널 × 식부자재로 별도 산출)

열 letter → 1-based 인덱스: K=11, S=19, Q=17, V=22, AG=33, AK=37
"""
from pathlib import Path
import openpyxl
from openpyxl.utils import column_index_from_string
from rich.console import Console

console = Console(highlight=False)

# 손익항목 → 정산내역서 열 letter
COUPANG_COLS = {
    "sales":        "K",   # 주문금액
    "ad_brokerage": "Q",   # 광고비(중개)
    "delivery":     "V",   # 배달비
    "card_fee":     "S",   # 카드/어플수수료
    "vat":          "AG",  # 부가세
    "ad_custom":    "AK",  # 광고비(맞춤)
}
# 비용 항목(절대값 양수로 변환) — 매출만 그대로
COST_KEYS = {"ad_brokerage", "delivery", "card_fee", "vat", "ad_custom"}


def _sum_column(ws, col_letter: str) -> float:
    """해당 열의 숫자 셀 전체 합 (헤더/문자 셀은 자동 무시)."""
    c = column_index_from_string(col_letter)
    total = 0.0
    for r in range(1, ws.max_row + 1):
        v = ws.cell(r, c).value
        if isinstance(v, (int, float)):
            total += v
    return total


def parse_coupang_settlement(path: Path) -> dict:
    """쿠팡이츠 정산내역서 → 손익 항목 dict."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    out = {}
    for key, col in COUPANG_COLS.items():
        val = _sum_column(ws, col)
        out[key] = abs(val) if key in COST_KEYS else val
    console.print(
        f"[green]  [쿠팡이츠] 파싱 완료 — 매출 {out['sales']:,.0f}원 "
        f"(중개 {out['ad_brokerage']:,.0f} / 배달 {out['delivery']:,.0f} / "
        f"카드 {out['card_fee']:,.0f} / 부가세 {out['vat']:,.0f} / 맞춤 {out['ad_custom']:,.0f})[/green]"
    )
    return out


if __name__ == "__main__":
    import sys, io
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
    p = Path(sys.argv[1]) if len(sys.argv) > 1 else None
    if p and p.exists():
        print(parse_coupang_settlement(p))
    else:
        print("사용: python step_coupang_parse.py <쿠팡이츠_정산내역서.xlsx>")
