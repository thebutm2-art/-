"""
S4. 비밀번호 해제 / 데이터 파싱
- 배민 정산명세서 엑셀을 암호(생년월일 6자리 / 법인번호 뒷 7자리)로 해제
- '상세' 시트의 3단 헤더(그룹/중간/세부)를 해석해 손익 항목으로 집계
- 컬럼 위치가 아닌 '헤더명 키워드' 기준으로 매핑 → 배민 양식 변경 내성 확보

[실제 배민 정산명세서 구조 — 2026-06 검증]
  시트: '요약'(카테고리 순액), '상세'(거래단위 명세)
  상세 헤더 3행:
    R3 그룹: (A) 주문중개 / (B) 배달 / (C) 그외 / (D) 기타 / (E) 부가세 / (F) 우리가게클릭 / (G) 배민오더 / (H) 입금금액
    R4 중간: 주문금액 / 중개이용료 / 고객할인비용 / 가게배달팁 / 배민클럽 할인비용 / 배달비 / 결제정산수수료 / 만나서결제금액 ...
    R5 세부: 바로결제주문금액 / 배민1중개이용료 / 메뉴할인 / ...
  R6~ : 데이터
"""
import io
import warnings
from pathlib import Path
from rich.console import Console
import msoffcrypto
import openpyxl

warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")

console = Console(highlight=False)

# 손익 항목 키 (step5 입력과 동일)
NUMERIC_KEYS = {
    "total_sales", "brokerage_fee", "payment_fee",
    "delivery_fee", "discount_burden", "ad_cost",
    "vat", "expected_settlement",
}


def _decrypt(path: Path, password: str) -> io.BytesIO:
    """암호화 엑셀 → 복호화된 BytesIO. 암호화 안 됐으면 원본 그대로."""
    with open(path, "rb") as f:
        office = msoffcrypto.OfficeFile(f)
        if not office.is_encrypted():
            return io.BytesIO(path.read_bytes())
        office.load_key(password=password)
        out = io.BytesIO()
        office.decrypt(out)
        out.seek(0)
        return out


GAGAE_DELIVERY_UNIT = 4400  # 가게배달(자체배달) 건당 배달비


def _classify_column(group: str, mid: str, sub: str) -> str | None:
    """그룹/중간/세부 헤더명으로 컬럼을 손익 항목 키에 분류 (사장님 신규 기준).

    매출       = 바로결제주문금액 + 가게배달팁(바로결제/만나서결제 배달팁) + 보정금액
    중개이용료  = 배민1 + 알뜰배달 + 가게배달(A부분중개) + 픽업 중개이용료 (전체)
    프로모션    = 고객할인비용 > 주문금액 즉시할인 (메뉴할인 제외)
    배달비      = 배민1 한집 + 알뜰배달 배달비 (+ 가게배달 건수×4400은 별도 가산)
    카드/어플   = 결제정산수수료 > 기본수수료(정률)  (우대수수료 제외)
    부가세      = 부가세
    광고비(우가클) = 우리가게클릭 이용요금 + 부가세
    """
    g = (group or "").replace(" ", "")
    m = (mid or "").replace(" ", "")
    s = (sub or "").replace(" ", "")

    # 보정금액 (그룹/중간/세부 어디든 '보정' 포함)
    if "보정" in g or "보정" in m or "보정" in s:
        return "total_sales"

    if "주문중개" in g:
        if "주문금액" in m:
            if "바로결제" in s:        return "total_sales"   # 만나서결제·부분환불 제외
            return None
        if "중개이용료" in m:          return "brokerage_fee"
        if "할인" in m:
            if "즉시할인" in s:        return "discount_burden"  # 메뉴할인 제외
            return None
    if "배달" in g and "주문중개" not in g:
        if "배달팁" in m:              return "total_sales"   # 가게배달팁 → 매출
        if "배달비" in m:              return "delivery_fee"  # 배민1한집 + 알뜰배달비
        return None                                          # 배민클럽 할인비용 등 제외
    if "그외" in g:
        if "결제" in m and "수수료" in m:
            if "기본수수료" in s:      return "payment_fee"   # 우대수수료 제외
            return None
    if "부가세" in g:              return "vat"
    if "우리가게클릭" in g:        return "ad_cost"
    if "입금금액" in g:            return "expected_settlement"
    return None


def _parse_baemin_detail(wb) -> dict:
    """배민 '상세' 시트를 손익 항목으로 집계."""
    ws = wb["상세"] if "상세" in wb.sheetnames else wb.active

    # 헤더 행 자동 탐지: '(A) 주문중개'가 들어있는 행을 그룹 헤더로
    group_row = None
    for r in range(1, 8):
        for c in range(1, ws.max_column + 1):
            v = ws.cell(r, c).value
            if v and "주문중개" in str(v):
                group_row = r
                break
        if group_row:
            break
    if not group_row:
        raise ValueError("상세 시트에서 헤더(주문중개)를 찾지 못함")

    mid_row  = group_row + 1
    sub_row  = group_row + 2
    data_row = group_row + 3  # 그룹/중간/세부 3행 뒤부터 데이터

    # 컬럼별 그룹/중간/세부 헤더 forward-fill
    col_class = {}
    cur_group = ""
    for c in range(1, ws.max_column + 1):
        g = ws.cell(group_row, c).value
        # 실제 그룹은 '(A)~(H)' 마커 형식만 인정 (예: '입금 금액' 소계열 오인 방지)
        if g and "(" in str(g) and ")" in str(g):
            cur_group = str(g)
        mid = ws.cell(mid_row, c).value
        sub = ws.cell(sub_row, c).value
        if mid:
            cur_mid = str(mid)
        else:
            prev = col_class.get(c - 1)
            cur_mid = prev[1] if (prev and prev[0] == cur_group) else ""
        col_class[c] = (cur_group, cur_mid, str(sub) if sub else "")

    result = {k: 0.0 for k in NUMERIC_KEYS}

    for c, (g, m, s) in col_class.items():
        key = _classify_column(g, m, s)
        if not key or key not in result:
            continue
        total = 0.0
        for r in range(data_row, ws.max_row + 1):
            v = ws.cell(r, c).value
            if isinstance(v, (int, float)):
                total += v
        result[key] += total

    # 비용 항목은 정산파일에서 음수 → 손익계산용 양수(절대값)로 변환
    for k in ("brokerage_fee", "payment_fee", "delivery_fee",
              "discount_burden", "ad_cost", "vat"):
        result[k] = abs(result[k])

    # 가게배달 건수 (E열 '주문유형/기타'에서 '가게배달' 행 수) — 참고/폴백용
    type_col = 5  # E열
    gagae = 0
    for r in range(data_row, ws.max_row + 1):
        v = ws.cell(r, type_col).value
        if v and "가게배달" in str(v):
            gagae += 1
    result["gagae_count_settlement"] = gagae
    return result


def parse_settlement_file(path: Path, store: dict,
                          gagae_count: int | None = None) -> dict | None:
    """점포 정산 첨부파일 파싱 → 손익 항목 dict. 실패 시 None.

    gagae_count: 가게배달 건수(배민셀프서비스 주문내역 기준). None이면 정산서 카운트 폴백.
                 배달비 += 가게배달건수 × 4400.
    """
    password = store.get("file_pw") or store.get("biz_no", "")
    try:
        dec = _decrypt(path, password)
        wb  = openpyxl.load_workbook(dec, data_only=True)

        if "상세" in wb.sheetnames or any("주문중개" in str(wb.active.cell(r, c).value)
                                          for r in range(1, 8)
                                          for c in range(1, min(wb.active.max_column + 1, 40))):
            data = _parse_baemin_detail(wb)
        else:
            console.print(f"[yellow]  [{store['name']}] 배민 표준 양식 아님 — 확인 필요[/yellow]")
            return None

        # 가게배달 배달비 = 건수 × 4400 (주문내역 기준 우선, 없으면 정산서 카운트)
        cnt = gagae_count if gagae_count is not None else data.get("gagae_count_settlement", 0)
        data["gagae_count"] = cnt
        data["gagae_delivery_fee"] = cnt * GAGAE_DELIVERY_UNIT
        data["delivery_fee"] += data["gagae_delivery_fee"]

        data["store_code"] = store["code"]
        data["store_name"] = store["name"]
        src = "주문내역" if gagae_count is not None else "정산서카운트"
        console.print(
            f"[green]  [{store['name']}] 파싱 완료 — 매출 {data['total_sales']:,.0f}원 "
            f"| 배달비 {data['delivery_fee']:,.0f}(가게배달 {cnt}건×4400={data['gagae_delivery_fee']:,.0f}, {src})[/green]"
        )
        return data

    except Exception as e:
        console.print(f"[red]  [{store['name']}] 파싱 실패 ({path.name}): {e}[/red]")
        return None


def parse_all(file_map: dict[str, list[Path]], stores: list[dict]) -> list[dict]:
    """전 점포 첨부 파싱. file_map: {점포코드: [Path,...]}"""
    store_idx = {s["code"]: s for s in stores}
    results   = []
    for code, paths in file_map.items():
        if not paths:
            console.print(f"[yellow]  [{code}] 첨부파일 없음 — 스킵[/yellow]")
            continue
        store = store_idx.get(code)
        if not store:
            continue
        for path in paths:
            data = parse_settlement_file(path, store)
            if data:
                results.append(data)
                break
    console.print(f"[bold green]파싱 완료: {len(results)}/{len(stores)}개 점포[/bold green]")
    return results
