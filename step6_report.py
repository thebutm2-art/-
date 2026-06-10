"""
S6. Excel 리포트 생성
- 점포별 손익계산서 (개별 시트)
- 전 점포 통합 집계표 (요약 시트)
- 실행 로그 (로그 시트)
"""
from datetime import datetime
from pathlib import Path
from rich.console import Console
import openpyxl
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side, numbers
)
from openpyxl.utils import get_column_letter
import config

console = Console()

# ── 스타일 상수 ──────────────────────────────────────────────
HEADER_FILL  = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT  = Font(color="FFFFFF", bold=True, size=10)
TITLE_FONT   = Font(bold=True, size=12)
PROFIT_FILL  = PatternFill("solid", fgColor="E2EFDA")
LOSS_FILL    = PatternFill("solid", fgColor="FCE4D6")
SUBTOTAL_FILL= PatternFill("solid", fgColor="BDD7EE")
NUM_FMT      = '#,##0'
PCT_FMT      = '0.0%'  # Excel 백분율 서식 (값 × 100)
THIN         = Side(style="thin")
BORDER       = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _cell(ws, row, col, value=None, *, bold=False, fill=None, num_fmt=None,
          font=None, align="left"):
    c = ws.cell(row=row, column=col, value=value)
    c.border    = BORDER
    c.alignment = Alignment(horizontal=align, vertical="center")
    if bold:
        c.font = Font(bold=True)
    if fill:
        c.fill = fill
    if num_fmt:
        c.number_format = num_fmt
    if font:
        c.font = font
    return c


def _write_store_sheet(wb: openpyxl.Workbook, pl: dict):
    """점포별 손익계산서 시트 생성"""
    title = f"{pl['store_code']}_{pl['store_name']}"[:31]  # Excel 시트명 최대 31자
    ws = wb.create_sheet(title=title)
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 18

    r = 1
    ws.merge_cells(f"A{r}:B{r}")
    c = ws.cell(r, 1, f"배민 손익계산서 — {pl['store_name']} ({pl['target_month']})")
    c.font      = TITLE_FONT
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[r].height = 20
    r += 1

    rows = [
        ("항목",              "금액(원)",          None,          True),
        ("(+) 총매출액",       pl["total_sales"],   NUM_FMT,       False),
        ("(-) 중개수수료",     pl["brokerage_fee"], NUM_FMT,       False),
        ("(-) 결제정산수수료", pl["payment_fee"],   NUM_FMT,       False),
        ("(-) 배달비(점주부담)",pl["delivery_fee"],  NUM_FMT,       False),
        ("(-) 할인·쿠폰 부담금",pl["discount_burden"],NUM_FMT,     False),
        ("(-) 광고비",         pl["ad_cost"],       NUM_FMT,       False),
        ("(=) 배달채널 기여이익",pl["contribution_profit"], NUM_FMT, "subtotal"),
        ("(-) 매출원가",       pl["cogs"],          NUM_FMT,       False),
        (f"    (원가율 {pl['cost_rate']*100:.1f}%)", "", None,      False),
        ("(-) 월 고정비",      pl["fixed_cost"],    NUM_FMT,       False),
        ("(=) 영업이익",       pl["operating_profit"], NUM_FMT,    "profit"),
        ("    영업이익률",      pl["operating_margin_pct"] / 100, PCT_FMT, False),
        ("",                   "",                  None,          False),
        ("[참고] 부가세(VAT)",  pl["vat"],           NUM_FMT,       False),
    ]

    for label, value, fmt, special in rows:
        if special is True:  # 헤더
            _cell(ws, r, 1, label, fill=HEADER_FILL, font=HEADER_FONT, align="center")
            _cell(ws, r, 2, value, fill=HEADER_FILL, font=HEADER_FONT, align="center")
        elif special == "subtotal":
            _cell(ws, r, 1, label, fill=SUBTOTAL_FILL, bold=True)
            _cell(ws, r, 2, value, fill=SUBTOTAL_FILL, bold=True, num_fmt=fmt, align="right")
        elif special == "profit":
            fill = PROFIT_FILL if (value or 0) >= 0 else LOSS_FILL
            _cell(ws, r, 1, label, fill=fill, bold=True)
            _cell(ws, r, 2, value, fill=fill, bold=True, num_fmt=fmt, align="right")
        else:
            _cell(ws, r, 1, label)
            _cell(ws, r, 2, value, num_fmt=fmt, align="right")
        r += 1


def _write_summary_sheet(wb: openpyxl.Workbook, pl_list: list[dict]):
    """통합 집계 요약 시트"""
    ws = wb.create_sheet(title="통합집계", index=0)
    cols = [
        ("점포코드", 10), ("점포명", 14), ("총매출액", 16),
        ("중개수수료", 14), ("결제수수료", 14), ("배달비", 14),
        ("할인부담금", 14), ("광고비", 12),
        ("기여이익", 16), ("매출원가", 14), ("고정비", 14),
        ("영업이익", 16), ("영업이익률", 12),
    ]
    for i, (name, width) in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
        c = ws.cell(1, i, name)
        c.fill      = HEADER_FILL
        c.font      = HEADER_FONT
        c.border    = BORDER
        c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 18

    # 영업이익 기준 내림차순 정렬 (이상치 파악 용이)
    sorted_pl = sorted(pl_list, key=lambda x: x["operating_profit"], reverse=True)

    for r, pl in enumerate(sorted_pl, 2):
        profit = pl["operating_profit"]
        fill   = PROFIT_FILL if profit >= 0 else LOSS_FILL
        row_data = [
            (pl["store_code"],          None,    "left"),
            (pl["store_name"],          None,    "left"),
            (pl["total_sales"],         NUM_FMT, "right"),
            (pl["brokerage_fee"],       NUM_FMT, "right"),
            (pl["payment_fee"],         NUM_FMT, "right"),
            (pl["delivery_fee"],        NUM_FMT, "right"),
            (pl["discount_burden"],     NUM_FMT, "right"),
            (pl["ad_cost"],             NUM_FMT, "right"),
            (pl["contribution_profit"], NUM_FMT, "right"),
            (pl["cogs"],                NUM_FMT, "right"),
            (pl["fixed_cost"],          NUM_FMT, "right"),
            (profit,                    NUM_FMT, "right"),
            (pl["operating_margin_pct"] / 100, PCT_FMT, "right"),
        ]
        for col, (val, fmt, align) in enumerate(row_data, 1):
            c = ws.cell(r, col, val)
            c.border    = BORDER
            c.alignment = Alignment(horizontal=align, vertical="center")
            if fmt:
                c.number_format = fmt
            if col >= 12:
                c.fill = fill

    # 합계 행
    total_row = len(sorted_pl) + 2
    ws.cell(total_row, 1, "합계").font = Font(bold=True)
    ws.cell(total_row, 1).border = BORDER
    num_cols = {3, 4, 5, 6, 7, 8, 9, 10, 11, 12}
    for col in range(2, 14):
        c = ws.cell(total_row, col)
        c.border    = BORDER
        c.fill      = SUBTOTAL_FILL
        c.font      = Font(bold=True)
        if col in num_cols:
            key_map = {3:"total_sales",4:"brokerage_fee",5:"payment_fee",
                       6:"delivery_fee",7:"discount_burden",8:"ad_cost",
                       9:"contribution_profit",10:"cogs",11:"fixed_cost",12:"operating_profit"}
            key = key_map.get(col)
            if key:
                c.value          = sum(p[key] for p in pl_list)
                c.number_format  = NUM_FMT
                c.alignment      = Alignment(horizontal="right")


def _write_log_sheet(wb: openpyxl.Workbook, log_entries: list[dict]):
    """실행 로그 시트"""
    ws = wb.create_sheet(title="실행로그")
    headers = ["점포코드", "점포명", "상태", "사유", "처리시각"]
    for i, h in enumerate(headers, 1):
        c = ws.cell(1, i, h)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.border = BORDER

    for r, entry in enumerate(log_entries, 2):
        for i, key in enumerate(["code", "name", "status", "reason", "ts"], 1):
            c = ws.cell(r, i, entry.get(key, ""))
            c.border = BORDER
            if key == "status":
                c.fill = PROFIT_FILL if entry.get("status") == "성공" else LOSS_FILL

    for col in ["A", "B", "C", "D", "E"]:
        ws.column_dimensions[col].width = 18


def generate_report(
    pl_list: list[dict],
    log_entries: list[dict],
    target_month: str,
    out_dir: Path | None = None,
) -> Path:
    """
    Excel 리포트 생성 → 파일 경로 반환
    """
    out_dir  = out_dir or config.OUTPUT_DIR
    filename = f"배민손익_{target_month}.xlsx"
    out_path = out_dir / filename

    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # 기본 Sheet 제거

    _write_summary_sheet(wb, pl_list)
    for pl in sorted(pl_list, key=lambda x: x["store_code"]):
        _write_store_sheet(wb, pl)
    _write_log_sheet(wb, log_entries)

    wb.save(out_path)
    console.print(f"\n[bold green]✓ 리포트 저장: {out_path}[/bold green]")
    return out_path
