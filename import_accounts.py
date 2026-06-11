"""
'배민 손익 계정 정보' 엑셀 → 점포마스터_v2.xlsx 변환

원본: 매장별 계정정보 시트 (C=매장명, D=배민아이디, E=배민비밀번호, 데이터 7행~)
생성: 점포마스터_v2.xlsx
  점포명 | 토더매장명 | 배민파트너명 | 파일암호 | 배민아이디 | 배민비밀번호 | 쿠팡아이디 | 쿠팡비밀번호 | 가게배달건수
  - 토더매장명: 기본값 = 점포명 (다르면 수정)
  - 배민파트너명/파일암호: 정산서 매칭·해제용 (계정파일에 없음 → 수기 입력 필요)
  - 가게배달건수: 월별로 배민셀프 가게통계에서 확인해 입력(선택)
"""
import sys, io
from pathlib import Path
import openpyxl

if sys.platform == "win32":
    try: sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
    except Exception: pass

BASE = Path(__file__).parent
COLS = ["점포명","토더매장명","배민파트너명","파일암호",
        "배민아이디","배민비밀번호","쿠팡아이디","쿠팡비밀번호","가게배달건수"]


def _clean(v):
    s = str(v or "").strip()
    return "" if s in ("계정 없음", "없음", "None", "nan") else s


def import_accounts(src: Path, out: Path | None = None) -> Path:
    out = out or (BASE / "점포마스터_v2.xlsx")
    wb = openpyxl.load_workbook(src, data_only=True)
    ws = wb["매장별 계정정보"] if "매장별 계정정보" in wb.sheetnames else wb.active

    # 헤더(매장명/아이디/비밀번호)가 있는 행 탐색
    hdr = None
    for r in range(1, 15):
        row = [str(ws.cell(r, c).value or "") for c in range(1, ws.max_column + 1)]
        if any("매장명" in x for x in row) and any("아이디" in x for x in row):
            hdr = r; break
    if not hdr:
        raise ValueError("계정정보 헤더(매장명/아이디)를 찾지 못함")

    # 그룹 헤더(배달의민족/쿠팡이츠) 행 = 헤더 바로 위
    grp = hdr - 1
    c_name = next((c for c in range(1, ws.max_column + 1)
                   if "매장명" in str(ws.cell(hdr, c).value or "")), None)

    # 그룹별 아이디/비번 열: 그룹 라벨 위치부터 오른쪽으로 아이디/비번 매칭
    def grp_cols(label):
        start = next((c for c in range(1, ws.max_column + 1)
                      if label in str(ws.cell(grp, c).value or "")), None)
        if not start:
            return None, None
        cid = cpw = None
        for c in range(start, ws.max_column + 1):
            h = str(ws.cell(hdr, c).value or "")
            if "아이디" in h and cid is None: cid = c
            elif "비밀번호" in h and cpw is None: cpw = c
            # 다음 그룹 시작 전까지만
            if c > start and str(ws.cell(grp, c).value or "").strip() and label not in str(ws.cell(grp, c).value):
                break
        return cid, cpw

    bm_id, bm_pw = grp_cols("배달의민족")
    cp_id, cp_pw = grp_cols("쿠팡")

    rows = []
    for r in range(hdr + 1, ws.max_row + 1):
        name = ws.cell(r, c_name).value
        if not name or not str(name).strip():
            continue
        rows.append({
            "점포명": str(name).strip(),
            "토더매장명": str(name).strip(),   # 기본값=점포명, 다르면 수정
            "배민파트너명": "",
            "파일암호": "",
            "배민아이디": _clean(ws.cell(r, bm_id).value) if bm_id else "",
            "배민비밀번호": _clean(ws.cell(r, bm_pw).value) if bm_pw else "",
            "쿠팡아이디": _clean(ws.cell(r, cp_id).value) if cp_id else "",
            "쿠팡비밀번호": _clean(ws.cell(r, cp_pw).value) if cp_pw else "",
            "가게배달건수": "",
        })

    nwb = openpyxl.Workbook(); nws = nwb.active; nws.title = "점포마스터"
    nws.append(COLS)
    for row in rows:
        nws.append([row[c] for c in COLS])
    nwb.save(out)
    print(f"[OK] 점포마스터 생성: {out}  ({len(rows)}개 매장)")
    print("  ※ 손익계산에 필요한 '파일암호'(생년월일6/법인번호뒤7), '배민파트너명'은 매장별로 채워주세요.")
    return out


if __name__ == "__main__":
    src = Path(sys.argv[1]) if len(sys.argv) > 1 else BASE / "계정정보_원본.xlsx"
    import_accounts(src)
